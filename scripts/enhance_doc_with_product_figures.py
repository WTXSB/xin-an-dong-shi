from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "心安动识_初赛作品说明文档.docx"
OUTPUT = ROOT / "心安动识_初赛作品说明文档_产品图示增强版.docx"
SCREENSHOTS = ROOT / "design-output" / "product-screenshots"

TEAL = RGBColor(43, 111, 99)
INK = RGBColor(54, 64, 62)
MUTED = RGBColor(91, 106, 102)

FIGURES = [
    ("相较于上述竞品", "01-product-home.png", "心安动识产品首页与核心服务入口",
     "产品首页将温柔感知、安心对话、情绪画像与正念片刻集中于同一服务入口，使“感知—理解—陪伴—行动”的产品闭环在首屏即可被清晰识别。"),
    ("图片感知、视频感知、摄像头感知三大核心检测闭环", "02-image-awareness.png", "图片温柔感知与用户自主授权界面",
     "图片感知在识别前明确告知用途，并将是否保存觉察记录、是否保留素材路径交由用户决定；模型结果被表达为可理解的情绪线索，而非诊断式标签。"),
    ("所有YOLO推理运算在浏览器本地完成", "03-video-awareness.png", "视频动态情绪线索分析界面",
     "视频感知面向连续动作与表情变化，呈现模型选择、识别阈值、处理进度和素材留存策略，将动态分析能力与可解释、可撤回的隐私控制结合。"),
    ("所有YOLO推理运算在浏览器本地完成", "04-camera-awareness.png", "摄像头实时身体与情绪信号感知界面",
     "摄像头模块仅在用户主动确认后开启，支持随时停止并整理本次线索；实时能力没有被包装成监控，而是被设计为可控、低压力的自我觉察工具。"),
    ("四大智能体形成", "06-safe-dialogue.png", "安心对话智能陪伴与隐私边界界面",
     "安心对话以陪伴式语言承接用户主动输入，并明确说明图片、摄像头画面与识别记录不会自动进入对话；这一区隔强化了心理支持场景中的知情、授权与数据最小化原则。"),
    ("四大智能体形成", "05-awareness-records.png", "觉察记录全流程沉淀与检索界面",
     "觉察记录统一承接图片、视频与摄像头三类感知结果，并提供分类、检索与删除入口，帮助用户回看状态变化，同时保留对个人数据的持续控制权。"),
    ("让每一次数据查看与功能使用都成为微型疗愈体验", "07-mind-spa.png", "心灵 SPA 真实支持资源导航界面",
     "心灵 SPA 在获得位置授权后连接附近真实支持资源，并将授权状态、定位范围与服务类别清晰呈现，把线上陪伴延伸至可触达的线下心理支持。"),
]


def style_run(run, size, color, bold=False):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")


def insert_after(anchor, paragraph):
    anchor._p.addnext(paragraph._p)
    return paragraph


def add_seq_field(run, number):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "SEQ Figure \\* ARABIC")
    result_run = OxmlElement("w:r")
    result_props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    result_props.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    result_props.append(size)
    result_run.append(result_props)
    text = OxmlElement("w:t")
    text.text = str(number)
    result_run.append(text)
    field.append(result_run)
    run._r.addnext(field)


def build_figure(doc, path, number, title, description):
    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(8)
    image_p.paragraph_format.space_after = Pt(3)
    image_p.paragraph_format.keep_with_next = True
    drawing = image_p.add_run().add_picture(str(path), width=Inches(6.15))
    drawing._inline.docPr.set("name", f"图 {number} {title}")
    drawing._inline.docPr.set("descr", description)

    caption_p = doc.add_paragraph(style="Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(0)
    caption_p.paragraph_format.space_after = Pt(4)
    caption_p.paragraph_format.keep_with_next = True
    lead = caption_p.add_run("图 ")
    style_run(lead, 9, TEAL, True)
    add_seq_field(lead, number)
    title_run = caption_p.add_run(f"  {title}")
    style_run(title_run, 9, INK, True)

    desc_p = doc.add_paragraph()
    desc_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    desc_p.paragraph_format.left_indent = Inches(0.18)
    desc_p.paragraph_format.right_indent = Inches(0.18)
    desc_p.paragraph_format.space_before = Pt(0)
    desc_p.paragraph_format.space_after = Pt(9)
    desc_p.paragraph_format.line_spacing = 1.2
    desc_p.paragraph_format.keep_together = True
    label = desc_p.add_run("功能说明：")
    style_run(label, 9.5, TEAL, True)
    body = desc_p.add_run(description)
    style_run(body, 9.5, MUTED)
    return image_p, caption_p, desc_p


def main():
    doc = Document(SOURCE)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace('实现"安小宁"智能体', '实现“安心对话”智能体')

    markers = {item[0] for item in FIGURES}
    anchors = {}
    for paragraph in doc.paragraphs:
        for marker in markers:
            if marker in paragraph.text:
                anchors[marker] = paragraph
    missing = sorted(markers - set(anchors))
    if missing:
        raise RuntimeError(f"Missing insertion anchors: {missing}")

    current_after = {}
    for number, (marker, filename, title, description) in enumerate(FIGURES, 1):
        anchor = current_after.get(marker, anchors[marker])
        image_path = SCREENSHOTS / filename
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        for part in build_figure(doc, image_path, number, title, description):
            anchor = insert_after(anchor, part)
        current_after[marker] = anchor

    props = doc.core_properties
    props.title = "心安动识初赛作品说明文档（产品图示增强版）"
    props.subject = "基于真实全栈网页截图的产品功能说明"
    props.keywords = "心安动识, 情绪识别, 心灵疗愈, 安心对话, YOLO"

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
